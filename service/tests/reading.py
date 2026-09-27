"""Helpers for tests that read files: synthetic BOQs, and waiting for the background reader."""

import io
import time
from pathlib import Path

import openpyxl
from PIL import Image, ImageDraw, ImageFont

FIXTURES = Path(__file__).parent / "fixtures"


def workbook() -> bytes:
    """A small made-up BOQ: an English sheet, an Arabic sheet and a hidden sheet."""
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Div.03"
    sheet["A1"], sheet["B1"] = "PROJECT:", "Synthetic Tower"
    for column, label in zip("ABCDEF", ["ITEM", "DESCRIPTION", "UNIT", "QTY", "RATE", "AMOUNT"], strict=True):
        sheet[f"{column}3"] = label
    rows = [
        (5, ["", "DIVISION 03 - CONCRETE"]),
        (6, ["", "Notes:"]),
        (7, ["", "* Contractor to allow for all testing"]),
        (8, ["3.1", "Cast-in-place concrete"]),
        (9, ["3.1.1", "Plain concrete grade C15 blinding", "m3", 86, 450, 38700]),
        (10, ["3.1.2", "Reinforced concrete to raft foundations", "m3", 1240.5, 1150.25, 1426885.13]),
        (11, ["3.1.3", "Formwork to sides of foundations", "m2", "1,250.00", "", ""]),
        (12, ["3.1.4", "Waterstops to construction joints", "L.M.", "", "", ""]),
        (13, ["3.1.5", "", "", 12, "", ""]),
        (14, ["", "Total carried to summary", "", "", "", 1465585.13]),
    ]
    for number, values in rows:
        for column, value in zip("ABCDEF", values, strict=False):
            if value != "":
                sheet[f"{column}{number}"] = value
    sheet.merge_cells("B1:D1")
    arabic = book.create_sheet("الملخص")
    arabic["A1"], arabic["B1"], arabic["C1"], arabic["D1"] = "البند", "الوصف", "الوحدة", "الكمية"
    arabic["A2"], arabic["B2"], arabic["C2"], arabic["D2"] = "1", "أعمال الحفر", "م3", 350
    hidden = book.create_sheet("Rates")
    hidden["A1"] = "internal rates"
    hidden.sheet_state = "hidden"
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def scan() -> bytes:
    """A page image with a small table drawn on it: no text layer, like a scanned BOQ page."""
    image = Image.new("RGB", (900, 400), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.load_default(size=28)
    for y, line in [
        (40, "ITEM   DESCRIPTION             UNIT   QTY"),
        (100, "1.1    Excavation in soil      m3     1240"),
    ]:
        draw.text((30, y), line, fill="black", font=font)
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def start(client, files: dict[str, bytes]) -> dict:
    response = client.post("/projects", files=[("files", (name, data)) for name, data in files.items()])
    assert response.status_code == 201, response.text
    return response.json()


def read_all(client, project_id: str, timeout: float = 30) -> dict:
    """The project once the background reader has finished with every file."""
    deadline = time.monotonic() + timeout
    while True:
        project = client.get(f"/projects/{project_id}").json()
        if all(s["status"] in ("read", "failed") for s in project["sources"]):
            return project
        assert time.monotonic() < deadline, f"reading did not finish: {project['sources']}"
        time.sleep(0.1)
