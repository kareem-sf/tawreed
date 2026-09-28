"""Publishing a revision: every workbook, values exactly as the source states them, all or nothing."""

import hashlib
import io
import json
import re
import zipfile
from decimal import Decimal

import openpyxl
import pytest
from reading import read_all, start, workbook

from tawreed import decisions, ledger, packages, publish
from tawreed.ledger.extract import SheetLayout
from tawreed.projects import Project
from tawreed.sources import Source

DIV03 = SheetLayout(first_row=4, code="A", description=["B"], unit="C", quantity="D", rate="E", amount="F")
ARABIC = SheetLayout(first_row=2, code="A", description=["B"], unit="C", quantity="D")
PLAN = ["Concrete works", "Formwork and joints", "Earthworks"]


def ready(client, files=None) -> str:
    """A project read, laid out, planned and placed, as the agent would leave it, with publishing requested."""
    project = read_all(client, start(client, files or {"Tower BOQ.xlsx": workbook()})["id"])
    home = client.app.state.home
    with client.app.state.sessions() as session:
        source = session.get(Source, project["sources"][0]["id"])
        ledger.lay_out_sheet(session, home, source, 1, DIV03, "agent")
        ledger.lay_out_sheet(session, home, source, 2, ARABIC, "agent")
        ledger.skip_pages(session, source, [3], "internal rates", "agent")
        me = session.get(Project, project["id"])
        concrete, formwork, earthworks = packages.apply_plan(
            session, me, [{"name": n, "scope": "", "reason": "", "keeps": []} for n in PLAN]
        )
        items = {i.ref: i for i in packages.active_items(session, me.id)}
        packages.place(session, me, [items[1], items[2]], concrete, "agent", "Concrete")
        packages.place(session, me, [items[3]], formwork, "agent")
        packages.place(session, me, [items[4]], formwork, "engineer", "The engineer chose this package.")
        packages.place(session, me, [items[5]], earthworks, "agent")
        decisions.raise_decision(session, me.id, "publish", {"summary": "Three packages."}, "agent")
        session.commit()
    return project["id"]


def approve(client, project_id):
    work = client.get(f"/projects/{project_id}/work").json()
    decision = next(d for d in work["decisions"] if d["kind"] == "publish")
    return client.post(f"/projects/{project_id}/decisions/{decision['id']}", json={"approve": True})


def folder(client, project_id, name="Rev 00"):
    return publish.revisions_folder(client.app.state.home, project_id) / name


def rows(sheet, first=1):
    return [list(row) for row in sheet.iter_rows(min_row=first, values_only=True)]


def computed(sheet, ref: str):
    """What Excel shows in a cell: its value, or the result of one of Tawreed's simple formulas."""
    value = sheet[ref].value
    if not isinstance(value, str) or not value.startswith("="):
        return value
    formula, book = value[1:], sheet.parent
    other = re.fullmatch(r"'((?:[^']|'')+)'!(.+)", formula)
    if other:  # on another sheet
        target = book[other[1].replace("''", "'")]
        inner = other[2]
        return computed(target, inner) if re.fullmatch(r"[A-Z]+\d+", inner) else _range_sum(target, inner)
    if match := re.fullmatch(r"([A-Z]+\d+)\*([A-Z]+\d+)", formula):
        return (computed(sheet, match[1]) or 0) * (computed(sheet, match[2]) or 0)
    if match := re.fullmatch(r"([A-Z]+\d+)-([A-Z]+\d+)", formula):
        return (computed(sheet, match[1]) or 0) - (computed(sheet, match[2]) or 0)
    if match := re.fullmatch(r"SUM\((.+)\)", formula):
        inner = re.fullmatch(r"'((?:[^']|'')+)'!(.+)", match[1])
        return _range_sum(book[inner[1].replace("''", "'")], inner[2]) if inner else _range_sum(sheet, match[1])
    raise AssertionError(f"not one of Tawreed's simple formulas: {value}")


def _range_sum(sheet, cells: str):
    first, last = cells.split(":")
    column, top, bottom = re.match(r"[A-Z]+", first)[0], int(first[len(re.match(r"[A-Z]+", first)[0]) :]), None
    bottom = int(last[len(column) :])
    values = [computed(sheet, f"{column}{r}") for r in range(top, bottom + 1)]
    return sum(v for v in values if isinstance(v, (int, float)))


def test_publishing_writes_every_workbook_with_the_values_as_the_source_states_them(client):
    project_id = ready(client)
    assert approve(client, project_id).status_code == 204

    revision = folder(client, project_id)
    written = sorted(str(p.relative_to(revision).as_posix()) for p in revision.rglob("*") if p.is_file())
    assert written == [
        "Coverage check - Rev 00.xlsx",
        "Decision log - Rev 00.xlsx",
        "Packages/01 Concrete works - Rev 00.xlsx",
        "Packages/02 Formwork and joints - Rev 00.xlsx",
        "Packages/03 Earthworks - Rev 00.xlsx",
        "Tower BOQ - Master - Rev 00.xlsx",
        "manifest.json",
    ]
    manifest = json.loads((revision / "manifest.json").read_text(encoding="utf-8"))
    for entry in manifest["files"]:
        assert hashlib.sha256((revision / entry["path"]).read_bytes()).hexdigest() == entry["sha256"]

    master = openpyxl.load_workbook(revision / "Tower BOQ - Master - Rev 00.xlsx")
    assert master.sheetnames == ["Cover", "Packages", "01 Concrete works", "02 Formwork and joints", "03 Earthworks"]
    concrete = master["01 Concrete works"]
    assert rows(concrete, 5)[:4] == [
        ["Item", "Description", "Unit", "Qty", "Rate", "Amount", "Source"],
        [None, "DIVISION 03 - CONCRETE › 3.1 Cast-in-place concrete", None, None, None, None, None],
        ["3.1.1", "Plain concrete grade C15 blinding", "m3", 86, 450, "=D7*E7", "Tower BOQ.xlsx › Div.03 › row 9"],
        [
            "3.1.2",
            "Reinforced concrete to raft foundations",
            "m3",
            1240.5,
            1150.25,
            1426885.13,  # the source rounded 1,426,885.125 to the cent, so its own figure stays
            "Tower BOQ.xlsx › Div.03 › row 10",
        ],
    ]
    assert computed(concrete, "F7") == 38700  # =86*450: the source's own amount
    total = rows(concrete, 10)[0]
    assert total[1] == "Total of the amounts above" and total[5] == "=SUM(F6:F8)"
    assert computed(concrete, "F10") == pytest.approx(1465585.13)
    cover, index = master["Cover"], master["Packages"]
    assert index["E2"].value == "='01 Concrete works'!F10" and computed(index, "E2") == pytest.approx(1465585.13)
    assert computed(cover, "B7") == pytest.approx(1465585.13)  # =SUM of the package index
    formwork = master["02 Formwork and joints"]
    quantity = formwork["D7"]
    assert (formwork["A7"].value, quantity.value, quantity.number_format) == ("3.1.3", 1250, "#,##0.00")  # "1,250.00"
    assert concrete["D8"].number_format == "#,##0.0"  # 1240.5: grouped, its own one decimal
    assert (formwork["A8"].value, formwork["C8"].value, formwork["D8"].value) == ("3.1.4", "L.M.", None)

    package = openpyxl.load_workbook(revision / "Packages/03 Earthworks - Rev 00.xlsx").active
    # Its only item has no amount, so there is no total: 0.00 would read as priced at nothing.
    assert rows(package, 6) == [["1", "أعمال الحفر", "م3", 350, None, None, "Tower BOQ.xlsx › الملخص › row 2"]]

    work = client.get(f"/projects/{project_id}/work").json()
    assert work["stage"] == "published" and work["published"]["name"] == "Rev 00"
    assert work["answered"][-1]["answer"]["revision"] == "Rev 00"
    assert client.get("/projects").json()[0]["revision"] == "Rev 00"


def test_the_coverage_check_and_decision_log_account_for_every_item(client):
    project_id = ready(client)
    approve(client, project_id)
    revision = folder(client, project_id)

    check = openpyxl.load_workbook(revision / "Coverage check - Rev 00.xlsx")
    summary = {row[0]: row[1] for row in rows(check["Summary"]) if row[0]}
    assert summary["Items in use"] == 5 and summary["Placed in exactly one package"] == 5
    assert summary["Not placed"] == 0 and summary["Difference"] == "=B4-B5"
    assert computed(check["Summary"], "B6") == 0
    assert Decimal(str(summary["Amounts of all packages"])) == Decimal("1465585.13")
    assert [r[0] for r in rows(check["Items"], 2)] == ["3.1.1", "3.1.2", "3.1.3", "3.1.4", "1"]

    log = openpyxl.load_workbook(revision / "Decision log - Rev 00.xlsx")
    placed_by = {r[0]: r[4] for r in rows(log["Placements"], 2)}
    assert placed_by == {
        "3.1.1": "Tawreed's agent",
        "3.1.2": "Tawreed's agent",
        "3.1.3": "Tawreed's agent",
        "3.1.4": "The engineer",
        "1": "Tawreed's agent",
    }
    how = [r[2] for r in rows(log["Reading"], 2)]
    assert how[0].startswith("row 4; Columns: code A, description B, unit C, quantity D, rate E, amount F")
    assert how[2] == "Set aside: internal rates"


def test_a_failed_publish_leaves_nothing_behind(client, monkeypatch):
    project_id = ready(client)

    def broken(book):
        raise OSError("disk full")

    monkeypatch.setattr(publish.workbooks, "coverage_check", broken)
    with pytest.raises(OSError):
        approve(client, project_id)
    root = publish.revisions_folder(client.app.state.home, project_id)
    assert list(root.iterdir()) == []  # no revision, no staging folder
    work = client.get(f"/projects/{project_id}/work").json()
    assert [d["kind"] for d in work["decisions"]] == ["publish"]  # still waiting for the engineer


def test_each_publish_is_a_new_revision_and_none_is_overwritten(client):
    project_id = ready(client)
    approve(client, project_id)
    first = (folder(client, project_id) / "manifest.json").read_bytes()
    with client.app.state.sessions() as session:
        decisions.raise_decision(session, project_id, "publish", {"summary": "Again."}, "agent")
        session.commit()
    approve(client, project_id)
    assert (folder(client, project_id) / "manifest.json").read_bytes() == first
    assert folder(client, project_id, "Rev 01").is_dir()
    assert [r["name"] for r in client.get(f"/projects/{project_id}/revisions").json()] == ["Rev 01", "Rev 00"]


def test_publishing_waits_until_every_item_is_placed(client):
    project_id = ready(client)
    with client.app.state.sessions() as session:
        me = session.get(Project, project_id)
        packages.unplace(session, me, [packages.active_items(session, project_id)[0]])
        session.commit()
    response = approve(client, project_id)
    assert response.status_code == 409 and response.json()["detail"]["code"] == "not_ready_to_publish"
    assert not publish.revisions_folder(client.app.state.home, project_id).joinpath("Rev 00").exists()


def test_an_arabic_project_gets_arabic_right_to_left_workbooks(client):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "الأعمال"
    rows_ = [("البند", "الوصف", "الوحدة", "الكمية"), ("1", "أعمال الحفر", "م3", 350), ("2", "خرسانة عادية", "م3", 86)]
    for row in rows_:
        sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    project = read_all(client, start(client, {"مشروع.xlsx": buffer.getvalue()})["id"])
    with client.app.state.sessions() as session:
        source = session.get(Source, project["sources"][0]["id"])
        ledger.lay_out_sheet(session, client.app.state.home, source, 1, ARABIC, "agent")
        me = session.get(Project, project["id"])
        [works] = packages.apply_plan(session, me, [{"name": "أعمال مدنية", "scope": "", "reason": "", "keeps": []}])
        packages.place(session, me, packages.active_items(session, me.id), works, "agent")
        decisions.raise_decision(session, me.id, "publish", {"summary": ""}, "agent")
        session.commit()
    approve(client, project["id"])

    master = openpyxl.load_workbook(folder(client, project["id"]) / "مشروع - Master - Rev 00.xlsx")
    sheet = master["01 أعمال مدنية"]
    assert sheet.sheet_view.rightToLeft is True
    assert [c.value for c in sheet[5]] == ["البند", "الوصف", "الوحدة", "الكمية", "السعر", "المبلغ", "المصدر"]
    assert master.sheetnames[:2] == ["الغلاف", "الحزم"]


def test_revisions_are_opened_and_exported(client, monkeypatch):
    project_id = ready(client)
    approve(client, project_id)
    opened = []
    monkeypatch.setattr(publish.os, "startfile", opened.append, raising=False)
    monkeypatch.setattr(publish.os, "name", "nt")

    assert client.post(f"/projects/{project_id}/revisions/0/open").status_code == 204
    assert opened == [folder(client, project_id)]
    assert client.post(f"/projects/{project_id}/revisions/7/open").json()["detail"]["code"] == "revision_not_found"

    export = client.get(f"/projects/{project_id}/revisions/0/export")
    assert export.headers["content-type"] == "application/zip"
    assert "Tower%20BOQ%20-%20Rev%2000.zip" in export.headers["content-disposition"]
    names = zipfile.ZipFile(io.BytesIO(export.content)).namelist()
    assert "Rev 00/manifest.json" in names and "Rev 00/Packages/03 Earthworks - Rev 00.xlsx" in names


def test_excels_float_noise_is_shown_to_the_cent_and_real_decimals_as_they_are():
    from tawreed.publish.workbooks import _number_format

    assert _number_format("261016.92999999993") == "#,##0.00"  # a formula's float: the value stays, the display is tidy
    assert _number_format("1250.00") == "#,##0.00"
    assert _number_format("12.125") == "#,##0.000"
    assert _number_format("86") == "#,##0"


def test_package_workbooks_can_leave_rates_out_for_suppliers_to_price(client):
    project_id = ready(client)
    work = client.get(f"/projects/{project_id}/work").json()
    decision = next(d for d in work["decisions"] if d["kind"] == "publish")
    body = {"approve": True, "prices": False}
    assert client.post(f"/projects/{project_id}/decisions/{decision['id']}", json=body).status_code == 204
    revision = folder(client, project_id)

    package = openpyxl.load_workbook(revision / "Packages/01 Concrete works - Rev 00.xlsx").active
    assert [row[3:6] for row in rows(package, 7)[:2]] == [[86, None, "=D7*E7"], [1240.5, None, "=D8*E8"]]
    assert package["F10"].value == "=SUM(F6:F8)"  # the supplier's rates price the package
    package["E7"].value, package["E8"].value = 10, 20  # as a supplier fills them in
    assert computed(package, "F10") == 86 * 10 + 1240.5 * 20

    master = openpyxl.load_workbook(revision / "Tower BOQ - Master - Rev 00.xlsx")
    assert master["01 Concrete works"]["E7"].value == 450  # the master keeps the rates
    assert json.loads((revision / "manifest.json").read_text(encoding="utf-8"))["prices"] is False
    assert client.get(f"/projects/{project_id}/revisions").json()[0]["prices"] is False
